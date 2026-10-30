"""Check source-word preservation, navigation and linked reference consistency."""
from pathlib import Path
from io import BytesIO
import difflib,hashlib,json,re,subprocess,zipfile
from lxml import etree
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
BASE='8f1a8ed61664f3ce379f77b01353d3bcef12375a'
DOC='transcript/Syntheticism-Defence-21-September-2026-Reading-Edition.docx'
norm=lambda s:re.sub(r'\s+',' ',s).strip()
def source(name):return subprocess.check_output(['git','show',BASE+':'+name],cwd=ROOT)
def spoken(doc):
 active=False;out=[]
 for p in doc.paragraphs:
  if p.style.name=='Heading 1'and p.text=='References':break
  if p.style.name=='Speaker':active=True
  elif p.style.name=='Normal'and active:
   p=etree.fromstring(etree.tostring(p._p))
   for h in p.xpath('.//w:hyperlink[@w:anchor]',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):h.getparent().remove(h)
   out.append(''.join(p.itertext()))
 return out

def check():
 original=Document(BytesIO(source(DOC)));new=Document(ROOT/DOC)
 before=' '.join(spoken(original));after=' '.join(spoken(new))
 old_tokens=before.split();new_tokens=after.split()
 from edit_transcript import normalise, CORRECTIONS
 expected=[];rec=1;seen={}
 for i,p in enumerate(original.paragraphs):
  if p.text=='Recording 2':rec=2
  if p.style.name=='Speaker':
   m=re.fullmatch(r'(.+?)\s{2,}(\d\d:\d\d:\d\d)',p.text);base=f'r{rec}-'+m[2].replace(':','');seen[base]=seen.get(base,0)+1;anchor=base+('-'+str(seen[base]) if seen[base]>1 else '')
   expected.append(normalise(original.paragraphs[i+1].text,anchor))
 from lecture_layout import script_paragraphs
 keys=list(seen)
 first=keys.index('r1-001543');last=keys.index('r1-004817')
 expected=expected[:first]+[p['text']for p in script_paragraphs()]+expected[last+1:]
 assert norm(' '.join(expected))==norm(after),'Reading text differs from defence script or corrected discussion'
 changes=CORRECTIONS['corrections']
 text=(ROOT/'transcript/defence-transcript.md').read_text();body=text.split('**Stefan · 00:12:46**',1)[1].split('\n## References',1)[0]
 body=re.sub(r'^#{1,6} .*$', '',body,flags=re.M);body=re.sub(r'^\*\*.*?\*\*$', '',body,flags=re.M);body=re.sub(r'^\*IMG_.*?\*$', '',body,flags=re.M)
 body=re.sub(r'<a id="[^"]+"></a>','',body);body=re.sub(r'\[\^ref\d+\]','',body);body=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',body)
 body=re.sub(r'(?<!\*)\*([^*\n]+)\*(?!\*)',r'\1',body)
 assert norm(body)==norm(after),'Markdown/Word speech differs'
 assert not re.search(r'\barticles? (one|two|three|four|five|six|seven)\b',body,re.I)
 assert text.splitlines()[0]=='# Syntheticism: How I Learned to Love Democracy'
 labels=list(re.finditer(r'^\*\*([^*]+?) · \d\d:\d\d:\d\d\*\*$',text,re.M));names=[m[1]for m in labels];assert len(names)==90
 # A qualified uncertain attribution remains separate from a secure one.
 for a,b in zip(labels,labels[1:]):
  if a[1]==b[1]:assert re.search(r'^##',text[a.end():b.start()],re.M),(a[1],b[1])
 anchors=re.findall(r'<a id="(r[12]-[^"]+)"></a>',text);assert len(anchors)==202 and len(set(anchors))==202
 nav=json.loads((ROOT/'transcript/navigation.json').read_text());assert anchors==[a['anchor']for a in nav['anchors']]
 with zipfile.ZipFile(ROOT/DOC)as z:xml=etree.fromstring(z.read('word/document.xml'))
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 bookmarks=xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns);assert len(set(bookmarks))==len(bookmarks)
 assert all(a.replace('-','_')in bookmarks for a in anchors)
 layout=json.loads((ROOT/'documentation/lecture-layout.json').read_text())
 lecture_ids=re.findall(r'<a id="(lecture-p\d+)"></a>',text)
 assert lecture_ids==[p['id']for p in layout['paragraphs']]
 assert [n for p in layout['paragraphs']for n in p['slides_after']]==list(range(1,55))
 assert 'at the summit of spectacle, <a id="r1-001645"></a>voilà.'in text
 assert '\n\nBelow the image'not in text or text.count('Below the image')==1
 paragraphs=[p for p in new.paragraphs if p.style.name=='Speaker'];assert len(paragraphs)==90
 urls=re.findall(r'https://computerlars\.github\.io/Syntheticism-How-I-Learned-to-Love-Democracy/dissertation/[^\s)]+#page=(\d+)',text);assert all(1<=int(n)<=336 for n in urls)
 pdf=PdfReader(ROOT/'dissertation/Syntheticism-How-I-Learned-to-Love-Democracy.pdf');assert len(pdf.pages)==336
 for a in nav['articles']:
  page=a['printed_pages'][0];plain=norm(pdf.pages[page].extract_text());assert f'Article {a["number"]}.'in plain
  if nav['pdf']['url']:assert a['url'].endswith('#page='+str(page+1))
 assert 'database does not read' in norm(pdf.pages[29].extract_text())
 assert 'Two kinds of autonomy' in norm(pdf.pages[54].extract_text())
 manifest={row['path']:row for row in json.loads((ROOT/'FILE-MANIFEST.json').read_text())['files']}
 for path in (['dissertation/Syntheticism-How-I-Learned-to-Love-Democracy.pdf']if nav['pdf']['url']else[])+['defence/Syntheticism-Defence-21-September-2026.pptx']:assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==manifest[path]['sha256']
 assert nav['anchors'][next(i for i,x in enumerate(nav['anchors'])if x['anchor']=='r2-004514')]['speaker']=='Asker'
 assert nav['anchors'][next(i for i,x in enumerate(nav['anchors'])if x['anchor']=='r1-004844')]['speaker']=='Stefan'
 paragraph=next(p for p in new.paragraphs if 'solely through practice' in p.text)
 assert any(run.text=='research' and run.italic for run in paragraph.runs)
 assert 'practice-based *research*' in text
 assert 'You write once, ideology is not a judgment' not in text
 assert 'Idiocy is not a judgment. It is a form comparable to ideology.' in text
 assert all(n['anchor'] in anchors for n in nav['passage_references'])
 assert 'IMG_' not in text
 old_readme=source('README.md').decode().replace('Text reproduced from the included PDF;','Text reproduced from the dissertation;').replace('Teksten er gengivet fra den vedlagte PDF;','Teksten er gengivet fra afhandlingen;');readme=(ROOT/'README.md').read_text();assert old_readme.split('## English abstract\n',1)[1]==readme.split('## English abstract\n',1)[1]
 result={'source_words':len(old_tokens),'documented_local_corrections':len(changes),'speaker_turns':len(names),'source_time_anchors':202,'word_markdown_speech_identical':True,'opening_lecture_matches_defence_script':True,'recorded_discussion_matches_correction_register':True,'published_binaries_match_manifest':True,'abstracts_unchanged':True,'pdf_links_checked':len(urls),'printed_page_offset':1}
 print(json.dumps(result,indent=2))

if __name__=='__main__':check()

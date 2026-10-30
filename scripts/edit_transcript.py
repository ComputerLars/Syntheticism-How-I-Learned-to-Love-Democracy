"""Rebuild the linked reading edition from the author's preserved edition.

The source edition remains in Git history. Local corrections and author-supplied
readings are enumerated in documentation/transcript-corrections.json.
"""
from pathlib import Path
from io import BytesIO
from copy import deepcopy
from datetime import datetime, timezone
import hashlib, json, re, subprocess
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import RGBColor

ROOT=Path(__file__).resolve().parents[1]
BASE='8f1a8ed61664f3ce379f77b01353d3bcef12375a'
DOC='transcript/Syntheticism-Defence-21-September-2026-Reading-Edition.docx'
TITLE='Syntheticism: How I Learned to Love Democracy'
PDF='https://computerlars.github.io/Syntheticism-How-I-Learned-to-Love-Democracy/dissertation/Syntheticism-How-I-Learned-to-Love-Democracy.pdf'
DATE='2026-10-21'
ARTICLES=[
 (1,'The Synthetic Party & Leader Lars: On Algorithmic Representation within Political Democracy','Asker Bryld Staunæs & Computer Lars',135,155),
 (2,'A Model for Proustian Decay: Estimating the Value of Intelligence in the Age of Unreason','Asker Bryld Staunæs & Computer Lars',157,177),
 (3,'Isaac Asimov’s Incubation of Political AI: An Immanent Critique?','Asker Bryld Staunæs',180,205),
 (4,'Deep Faking in a Flat Reality?','Asker Bryld Staunæs & Maja Bak Herrie',207,233),
 (5,'Scripting the Spectacle: A Theory Tragedy','Asker Bryld Staunæs',235,265),
 (6,'Synthetic Chamber: Agentic Mediation in Representative Democracy','Michal Malý & Asker Bryld Staunæs',268,294),
 (7,'Democracy as a Governance Algorithm: A Constraint Hierarchy for the AI Society','Asker Bryld Staunæs',296,324),
]
BY_NUMBER={a[0]:a for a in ARTICLES}
PUBLICATIONS=json.loads((ROOT/'documentation/article-publications.json').read_text())
# Bibliographic locations, not replacements for the spoken formulations.
NOTES=[
 ('r1-002322','Geometry of the dissertation: Figure N.1, p. 14; directional variations, pp. 15–16.',[14,15,16]),
 ('r1-002821','Article 1: emotion analysis, figure and discussion, pp. 138–139.',[138,139]),
 ('r1-002909','Article 2: Proust quotation, p. 160; Figure 2.2, p. 161.',[160,161]),
 ('r1-003004','Proust book photograph and permutation: Figure N.4, p. 31; analysis and photographed order, pp. 32–35.',[31,32,35]),
 ('r1-003129','Guardian composite: Figure N.3, p. 27; synthesis discussion, pp. 28–29. The complete article is Article 4, pp. 207–233.',[27,28,29,207]),
 ('r1-003624','Wiktoria–Lars exchange, pp. 106–107; social and technosocial sculpture, pp. 86–94.',[106,107,86,94]),
 ('r1-004213','Synthetic Summit Simulator and proposed Synthetic Chamber: Article 6, pp. 268–294; constitutional workflow and uptake, pp. 279–285.',[268,279,285]),
 ('r1-004021','Opening formulation of democracy’s disproportion: printed p. 6.',[6]),
 ('r1-004120','Abstract formulation of democracy’s disproportion: printed p. 327.',[327]),
 ('r1-004817','The closing thesis formulation appears in the English abstract, p. 330. The spoken ending is the present defence formulation.',[330]),
 ('r1-004415','Collective Resolution and democratic specification: synthesis, pp. 107–109; Article 7, pp. 296–324.',[107,109,296]),
 ('r2-002017','Jan-Werner Müller’s reading of Asimov’s electoral computer and its limits: Article 3, pp. 197–199.',[197,198,199]),
 ('r2-003419','Emily Bender’s objection to The Synthetic Party is quoted and discussed on p. 85.',[85]),
 ('r2-005241','Stiegler’s idiotextual geometry: Article 5, pp. 245–247.',[245,246,247]),
 ('r2-010436','Retro-futurist scenography, p. 69. The earlier formulation is “consumable futurisms”, p. 58.',[69,58]),
]
WEB_SOURCES={
 'r1-011049':('Alexander Bogdanov, Red Star (Красная звезда), Saint Petersburg, 1908. First-edition catalogue, National Electronic Library / Russian State Library.','https://rusneb.ru/catalog/000199_000009_003746695/'),
 'r1-003416':('Synthetic Summit, Kunsthal Aarhus; exhibition and AI World Congress, 2025.','https://www.kunsthalaarhus.dk/en/Events/Opening-Of-Synthetic-Summit'),
 'r1-003524':('AI: Authority of Art, Ujazdowski Castle Centre for Contemporary Art, Warsaw, 18 July–21 September 2025.','https://u-jazdowski.pl/en/programme/exhibitions/ai-authority-of-art'),
 'r2-002758':('The Swedish event was Partikonvent & Valvaka – AI-partiet, Inkonst, Malmö, 9–13 September 2026.','https://inkonst.com/sv/events/partikonvent-and-valvaka-ai-partiet'),
 'r2-000834':('Joë Bousquet (1897–1950), a French poet, is discussed in Gilles Deleuze, The Logic of Sense, “Twenty-First Series of the Event”, pp. 148–153, especially p. 148. Translated by Mark Lester with Charles Stivale; edited by Constantin V. Boundas. New York: Columbia University Press, 1990. Deleuze quotes Bousquet: “My wound existed before me, I was born to embody it.”','https://archive.org/details/logicofsense0000dele/page/148/mode/2up'),
}
NOTES += [(anchor,text,[]) for anchor,(text,url) in WEB_SOURCES.items()]
NOTE_MAP={key:(i,text,pages)for i,(key,text,pages)in enumerate(NOTES,1)}

def source(path):
 return subprocess.check_output(['git','show',BASE+':'+path],cwd=ROOT)

def page_url(n):return PDF+'#page='+str(n+1)if json.loads((ROOT/'metadata.json').read_text())['dissertation'].get('publicly_available',True)else'#references'
def unbroken_range(text):return re.sub(r'(?<=\d)–(?=\d)','\u2060–\u2060',text)

CORRECTIONS=json.loads((ROOT/'documentation/transcript-corrections.json').read_text())

def normalise(text,anchor=None):
 for edit in CORRECTIONS['corrections']:
  if edit['anchor']==anchor:
   assert text.count(edit['before'])==1,(anchor,edit['before'])
   text=text.replace(edit['before'],edit['after'])
 number=r'(?:one|two|three|four|five|six|seven|[1-7])'
 separator=r'(?:\s+(?:and|through|to)\s+|,\s*|\s*[–-]\s*)'
 text=re.sub(r'\barticles?\s+'+number+r'\b(?:'+separator+number+r'\b)*',article_spelling,text,flags=re.I)
 for n,word in enumerate(['one','two','three'],1):text=re.sub(r'\b(research question) '+word+r'\b',lambda m:m[1]+' '+str(n),text,flags=re.I)
 return text

def article_spelling(m):
 text=re.sub(r'^articles?','Articles'if m[0].lower().startswith('articles')else'Article',m[0],flags=re.I)
 for i,word in enumerate(['one','two','three','four','five','six','seven'],1):text=re.sub(r'\b'+word+r'\b',str(i),text,flags=re.I)
 return text

def spans(text):
 out=[]
 pattern=r'\bArticles?\s+[1-7](?:(?:\s+(?:and|through|to)\s+|,\s*|\s*[–-]\s*)[1-7])*'
 for m in re.finditer(pattern,text):
  nums=list(re.finditer(r'[1-7]',m[0]))
  for j,num in enumerate(nums):
   a=m.start()+(0 if j==0 else num.start());b=m.start()+num.end();n=int(num[0]);article=BY_NUMBER[n]
   out.append((a,b,page_url(article[3]),f'Article {n}: {article[1]}; printed pp. {article[3]}–{article[4]}'))
 for i,word in enumerate(['first','second','third','fourth','fifth','sixth','seventh'],1):
  for m in re.finditer(r'\b'+word+r' (?:research )?article\b',text,re.I):out.append((m.start(),m.end(),page_url(BY_NUMBER[i][3]),'Article '+str(i)))
 for m in re.finditer(r'\bpage (29|54)\b',text):out.append((m.start(),m.end(),page_url(int(m[1])),f'Printed p. {m[1]}; PDF page {int(m[1])+1}'))
 for m in re.finditer(r'\bresearch question ([123])\b',text,re.I):out.append((m.start(),m.end(),page_url({1:40,2:63,3:81}[int(m[1])]),'Extended synthesis · RQ'+m[1]))
 return sorted(out)

def page_spans(text):
 return [(m.start(),m.end(),page_url(int(m[1])),'Printed folio '+m[1])for m in re.finditer(r'\bpp?\. (\d+)(?:[–-]\d+)?',text)]

def linked_md(text,span_fn=spans,anchor=None):
 cursor=0;out=''
 for a,b,url,title in span_fn(text):out+=text[cursor:a]+'['+text[a:b]+']('+url+')';cursor=b
 out+=text[cursor:]
 for rule in CORRECTIONS.get('emphasis',[]):
  if rule['anchor']==anchor and rule['phrase'] in out:
   assert out.count(rule['phrase'])==1
   out=out.replace(rule['phrase'],rule['phrase'].replace(rule['text'],'*'+rule['text']+'*'))
 return out

bookmark_id=0
def bookmark(p,name):
 global bookmark_id
 bookmark_id+=1;a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bookmark_id));a.set(qn('w:name'),name.replace('-','_'));b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bookmark_id));p._p.append(a);p._p.append(b)

def hyperlink(p,text,url=None,anchor=None,tip=None,superscript=False):
 h=OxmlElement('w:hyperlink')
 if url:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 if anchor:h.set(qn('w:anchor'),anchor)
 if tip:h.set(qn('w:tooltip'),tip)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');rp.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u)
 if superscript:
  v=OxmlElement('w:vertAlign');v.set(qn('w:val'),'superscript');rp.append(v)
 r.append(rp);t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=unbroken_range(text);r.append(t);h.append(r);p._p.append(h)

def plain_word(p,text,anchor):
 rule=next((r for r in CORRECTIONS.get('emphasis',[])if r['anchor']==anchor and r['phrase']in text),None)
 if not rule:
  p.add_run(text)
  return
 start=text.index(rule['phrase'])+rule['phrase'].index(rule['text']);end=start+len(rule['text'])
 p.add_run(text[:start]);p.add_run(text[start:end]).italic=True;p.add_run(text[end:])

def linked_word(p,text,span_fn=spans,anchor=None):
 cursor=0
 for a,b,url,title in span_fn(text):
  plain_word(p,text[cursor:a],anchor);hyperlink(p,text[a:b],url=url,tip=title);cursor=b
 plain_word(p,text[cursor:],anchor)

def continuation(before,after):
 # Join only source paragraph breaks that divide the same sentence.
 return bool(before and after and (after[:1].islower() or not re.search(r'[.!?…][”’"\']?$',before)))

def build():
 original=source(DOC);doc=Document(BytesIO(original));ps=list(doc.paragraphs);events=[];recording=1;last=None;merged=0;joins=0;anchors=[];expected=[];used={}
 for i,p in enumerate(ps):
  style=p.style.name
  if style=='Speaker':
   m=re.fullmatch(r'(.+?)\s{2,}(\d\d:\d\d:\d\d)',p.text);assert m,p.text
   speaker,time=m.groups();base=f'r{recording}-'+time.replace(':','');used[base]=used.get(base,0)+1;anchor=base+('-'+str(used[base])if used[base]>1 else'');text=normalise(ps[i+1].text,anchor);speaker=CORRECTIONS['speaker_corrections'].get(anchor,speaker)
   entry={'speaker':speaker,'time':time,'anchor':anchor,'recording':recording,'text':text,'original_paragraph':i+1}
   expected.append(text);anchors.append({k:v for k,v in entry.items()if k!='text'})
   if not text:
    events.append({'type':'omission','entry':entry});continue
   if last==speaker:merged+=1
   else:events.append({'type':'speaker','paragraph':p,'entry':entry})
   if last==speaker and events[-1]['type']=='speech' and continuation(events[-1]['parts'][-1]['text'],text):events[-1]['parts'].append(entry);joins+=1
   else:events.append({'type':'speech','paragraph':ps[i+1],'parts':[entry]})
   last=speaker
  elif i and ps[i-1].style.name=='Speaker':continue
  else:
   if p.text=='Recording 2':recording=2
   if style.startswith('Heading'):last=None
   events.append({'type':'metadata','paragraph':p,'original_index':i})
 assert len(anchors)==202
 from lecture_layout import reparagraph,script_paragraphs
 events=reparagraph(events)
 first=next(i for i,a in enumerate(anchors)if a['anchor']=='r1-001543')
 last=next(i for i,a in enumerate(anchors)if a['anchor']=='r1-004817')
 expected=expected[:first]+[p['text']for p in script_paragraphs()]+expected[last+1:]
 # Reuse the author's styles, margins, headers and footers.
 for child in list(doc._element.body):
  if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
 md=[];spoken=[]
 for event in events:
  kind=event['type']
  if kind=='metadata':
   old=event['paragraph'];index=event['original_index'];text=old.text;style=old.style.name
   if style=='Record' and 'IMG_' in text:continue
   if index==0:
    doc.add_paragraph(TITLE,'Title');md.append('# '+TITLE)
   elif index==1:
    doc.add_paragraph('Public defence · 21 September 2026 · Aarhus University\nReading edition','Subtitle');md.append('*Public defence · 21 September 2026 · Aarhus University · Reading edition*')
   elif index==2:
    p=doc.add_paragraph('Asker Bryld Staunæs\nKasernen, Lille Sal, Aarhus')
    note='The opening lecture follows the author’s defence script; its timestamps locate corresponding recorded passages. The discussions are edited from the recordings. Timestamps restart for recording 2. References use printed folios; the PDF page is one higher.'
    doc.add_paragraph(note,'Record');md.append('Asker Bryld Staunæs  \nKasernen, Lille Sal, Aarhus');md.append(note)
    md.append('[Opening remarks](#opening-remarks) · [Opening lecture](#opening-lecture) · [Jan Bäcklund](#discussion-with-jan-bäcklund) · [Sine Nørholm Just](#discussion-with-sine-nørholm-just) · [Jussi Parikka](#discussion-with-jussi-parikka) · [References](#references)')
    doc.add_paragraph('Recording 1','Heading 1');md.append('## Recording 1')
   else:
    doc._element.body.insert(-1,deepcopy(old._p))
    if style.startswith('Heading'):md.append(('#'* (2 if style=='Heading 1'else 3))+' '+text)
    elif style=='Record':md.append('*'+text+'*')
    elif text:md.append(text)
  elif kind=='omission':
   bookmark(doc.paragraphs[-1],event['entry']['anchor']);md.append('<a id="'+event['entry']['anchor']+'"></a>');spoken.append('')
  elif kind=='speaker':
   e=event['entry'];q=doc.add_paragraph(e['speaker']+'   '+e['time'],'Speaker');md.append('**'+e['speaker']+' · '+e['time']+'**')
  else:
   p=doc.add_paragraph(style=event['paragraph'].style.name);pieces=[]
   if event.get('paragraph_id'):
    bookmark(p,event['paragraph_id']);pieces.append('<a id="'+event['paragraph_id']+'"></a>\n\n')
   for j,part in enumerate(event['parts']):
    if j:p.add_run(' ');pieces.append(' ')
    if part.get('emit_anchor',True):
     bookmark(p,part['anchor']);pieces.append('<a id="'+part['anchor']+'"></a>')
    if j==0 and not event.get('paragraph_id'):pieces.append('\n\n')
    linked_word(p,part['text'],anchor=part['anchor']);pieces.append(linked_md(part['text'],anchor=part['anchor']))
    refs=([part['anchor']]if part.get('emit_reference',True)else[])+part.get('references',[])
    for reference in refs:
     if reference in NOTE_MAP:
      n,_,_=NOTE_MAP[reference];hyperlink(p,'['+str(n)+']',anchor='ref_'+str(n),superscript=True);pieces.append('[^ref'+str(n)+']')
    spoken.append(part['text'])
   md.append(''.join(pieces))
 assert re.sub(r'\s+',' ',' '.join(expected)).strip()==re.sub(r'\s+',' ',' '.join(spoken)).strip()
 doc.add_paragraph('References','Heading 1');md.append('## References')
 intro='Article, figure and passage references to the dissertation and related works.'
 doc.add_paragraph(intro);md.append(intro)
 doc.add_paragraph('Articles','Heading 2');md.append('### Articles')
 for n,title,authors,first,last in ARTICLES:
  stop=''if title.endswith(('.', '?', '!'))else'.'
  p=doc.add_paragraph();hyperlink(p,f'Article {n}',url=page_url(first));p.add_run(f'. {authors}. ');p.add_run(title).italic=True;p.add_run(unbroken_range(stop+' '+PUBLICATIONS[str(n)]+f' In dissertation: pp. {first}–{last}.'))
  md.append(f'- [Article {n}]('+page_url(first)+f'). {authors}. *{title}*'+stop+' '+PUBLICATIONS[str(n)]+f' In dissertation: pp. {first}–{last}.')
 doc.add_paragraph('Passage references','Heading 2');md.append('### Passage references')
 for i,(anchor,text,pages)in enumerate(NOTES,1):
  span_fn=page_spans if anchor!='r2-000834'else lambda t:[(a,b,WEB_SOURCES[anchor][1],'Deleuze, The Logic of Sense, p. 148')for a,b,_,_ in page_spans(t)]
  p=doc.add_paragraph();bookmark(p,'ref_'+str(i));p.add_run(str(i)+'. ');linked_word(p,text,span_fn)
  extra=''
  if anchor in WEB_SOURCES and WEB_SOURCES[anchor][1]:
   url=WEB_SOURCES[anchor][1];p.add_run(' ');hyperlink(p,'Source',url=url);extra=' [Source]('+url+').'
   if anchor=='r1-003524':
    p.add_run(' Linz: ');hyperlink(p,'KI-DIPFIES',url='https://memphismemph.is/program/ki-dipfies');p.add_run(' · ');hyperlink(p,'KI-DIPFEL GIPFELTREFFEN',url='https://memphismemph.is/program/ki-dipfel-gipfeltreffen');extra+=' Linz: [KI-DIPFIES](https://memphismemph.is/program/ki-dipfies) · [KI-DIPFEL GIPFELTREFFEN](https://memphismemph.is/program/ki-dipfel-gipfeltreffen).'
  md.append('[^ref'+str(i)+']: '+linked_md(text,span_fn)+extra)
 for name in ['Title','Subtitle','Heading 1','Heading 2','Speaker','Record']:
  doc.styles[name].font.color.rgb=RGBColor(0,0,0)
 doc.core_properties.author='Asker Bryld Staunæs';doc.core_properties.last_modified_by='Asker Bryld Staunæs';doc.core_properties.title=TITLE;doc.core_properties.subject='Public PhD defence, 21 September 2026; linked reading edition';doc.core_properties.modified=datetime(2026,10,21,0,0,tzinfo=timezone.utc)
 doc.save(ROOT/DOC);(ROOT/'transcript/defence-transcript.md').write_text('\n\n'.join(md)+'\n')
 navigation={'edition_date':DATE,'source_edition_commit':BASE,'source_docx_sha256':hashlib.sha256(original).hexdigest(),'source_timestamps':len(anchors),'speaker_turns':sum(e['type']=='speaker'for e in events),'omitted_interjections':[e['entry']['anchor']for e in events if e['type']=='omission'],'removed_repeated_speaker_labels':merged,'sentence_continuations_joined':joins,'anchors':anchors,'pdf':{'url':PDF,'printed_folio_offset':1},'articles':[{'number':n,'title':title,'authors':authors,'printed_pages':[first,last],'url':page_url(first)}for n,title,authors,first,last in ARTICLES],'passage_references':[{'number':i,'anchor':anchor,'reference':text,'printed_pages':pages}for i,(anchor,text,pages)in enumerate(NOTES,1)]}
 navigation['opening_lecture']={'basis':'authorial defence script','source':'defence/Syntheticism-Defence-21-September-2026.pptx','time_links':'corresponding recorded passages'}
 (ROOT/'transcript/navigation.json').write_text(json.dumps(navigation,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({k:navigation[k]for k in ['source_timestamps','speaker_turns','removed_repeated_speaker_labels','sentence_continuations_joined']},indent=2))

if __name__=='__main__':build()
